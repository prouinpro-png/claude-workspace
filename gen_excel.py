import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ── Couleurs ──────────────────────────────────────────────────────────────────
BLUE_FR   = "002395"   # bleu France
RED_FR    = "ED2939"   # rouge France
RED_ES    = "AA151B"   # rouge Espagne
YELLOW_ES = "F1BF00"   # jaune Espagne
HEADER_FG = "FFFFFF"
GREY_ROW  = "F2F2F2"
GREEN_OK  = "C6EFCE"
RED_WARN  = "FFC7CE"
ORANGE    = "FFEB9C"
BLUE_LIGHT= "DDEEFF"

def make_border():
    thin = Side(style="thin", color="CCCCCC")
    return Border(left=thin, right=thin, top=thin, bottom=thin)

def header_cell(ws, row, col, value, bg_color, fg_color="FFFFFF", bold=True, size=11):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(bold=bold, color=fg_color, size=size)
    c.fill = PatternFill("solid", fgColor=bg_color)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = make_border()
    return c

def data_cell(ws, row, col, value, bg_color=None, bold=False, wrap=True, align="left"):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(bold=bold, size=10)
    if bg_color:
        c.fill = PatternFill("solid", fgColor=bg_color)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    c.border = make_border()
    return c

# ═══════════════════════════════════════════════════════════════════════════════
# FEUILLE FRANCE
# ═══════════════════════════════════════════════════════════════════════════════
ws_fr = wb.active
ws_fr.title = "🇫🇷 France"
ws_fr.freeze_panes = "A3"

# Titre principal
ws_fr.merge_cells("A1:G1")
t = ws_fr["A1"]
t.value = "SUIVI COMMUNIQUÉ DE PRESSE — ISTÒRIA LAB — FRANCE"
t.font = Font(bold=True, color=HEADER_FG, size=13)
t.fill = PatternFill("solid", fgColor=BLUE_FR)
t.alignment = Alignment(horizontal="center", vertical="center")
ws_fr.row_dimensions[1].height = 28

# En-têtes colonnes
cols = ["Support", "Contact", "Email", "Envoi CP", "1ère relance", "2e relance", "Statut / Notes"]
for i, col in enumerate(cols, 1):
    header_cell(ws_fr, 2, i, col, RED_FR)
ws_fr.row_dimensions[2].height = 22

# Données France
france_data = [
    ("Le Film Français",            "Florian Krieg",                       "florian.krieg@lefilmfrancais.com",                     "01/04/2026", "07/04/2026", "—",          "En attente"),
    ("Satellifacts",                "Damien / Emmanuelle",                 "dc@satellifacts.com / em@satellifacts.com",             "01/04/2026", "—",          "—",          "En attente"),
    ("Cineuropa",                   "(équipe)",                            "laporta / gonzalez / boas @cineuropa.org",              "01/04/2026", "09/04/2026", "—",          "En attente"),
    ("Actu.fr / Le Métropolitain",  "Maxence Dourlen",                     "maxence.dourlen@actu.fr",                              "01/04/2026", "—",          "—",          "✅ Répondu — transféré vers Actu Toulouse"),
    ("France Télévisions",          "Coralie Pierre",                      "coralie.pierre@francetv.fr",                           "01/04/2026", "—",          "—",          "En attente"),
    ("Digivision",                  "Adrien",                              "a.nos@digivision.fr",                                  "01/04/2026", "—",          "—",          "En attente"),
    ("Midi Libre",                  "Hicham Tissaoui / Kenzo",             "htissaoui@midilibre.com / kenzofsr@gmail.com",          "01/04/2026", "—",          "—",          "En attente"),
    ("Hérault Tribune",             "Mégane Fernandez",                    "megane.fernandez@herault-tribune.com",                  "01/04/2026", "09/04/2026", "—",          "En attente"),
    ("Midi Libre",                  "Jérémy Bernede",                      "jbernede@midilibre.com",                               "02/04/2026", "07/04/2026", "10/04/2026", "En attente"),
    ("France 3 / France TV Tlse",   "Franck Omer",                         "franck.omer@francetv.fr / redactiontoulouse@francetv.tv","07/04/2026", "09/04/2026", "—",          "En attente"),
    ("Radio France",                "Sandrine Morin / Bénédicte Dupont",   "sandrine.morin@radiofrance.com / benedicte.dupont@radiofrance.com","07/04/2026","09/04/2026","—", "En attente"),
    ("RFM Montpellier",             "Romain",                              "redac-montpellier@rfm.fr",                             "07/04/2026", "09/04/2026", "—",          "En attente"),
    ("20 Minutes Toulouse",         "Lucie Tollon",                        "ltollon@20minutes.fr / toulouse@20minutes.fr",          "07/04/2026", "09/04/2026", "—",          "En attente"),
    ("La Dépêche du Midi",          "Vincent Dulong",                      "vincent.dulong@ladepechenews.fr",                      "07/04/2026", "09/04/2026", "—",          "En attente"),
    ("Éditions 31",                 "Jean",                                "redaction31@editions31.com",                           "07/04/2026", "09/04/2026", "—",          "En attente"),
    ("Snobinart",                   "Thibault Loucheux",                   "thibault.loucheux@snobinart.fr",                       "07/04/2026", "09/04/2026", "—",          "En attente"),
    ("Parcours des Arts",           "Yann",                                "contact@parcoursdesarts.com",                          "07/04/2026", "09/04/2026", "—",          "En attente"),
    ("Action Rédaction",            "Sévérine Martin / Alexandra Foissac", "severine.martin@action-redaction.fr / alexandra.foissac@wanadoo.fr","07/04/2026","09/04/2026","—","En attente"),
    ("Midi Libre",                  "Thierry",                             "tmbom@midilibre.com",                                  "09/04/2026", "—",          "—",          "En attente"),
    ("Grizette",                    "Laurent",                             "redaction@grizette.com",                               "09/04/2026", "—",          "—",          "En attente"),
    ("Claap.fr",                    "Thierry",                             "redaction@claap.fr",                                   "09/04/2026", "—",          "—",          "En attente"),
    ("Artistes de France",          "Anne Devailly",                       "a.devailly@artistes-de-france.com",                    "09/04/2026", "—",          "—",          "En attente"),
    ("Culture 31",                  "Bruno",                               "contact@culture31.com",                                "09/04/2026", "—",          "—",          "En attente"),
    ("La Région Occitanie",         "Nicolas Hubert",                      "nicolas.hubert@laregion.fr",                           "09/04/2026", "—",          "—",          "📵 Absent jusqu'au 13/04"),
    ("La Région Occitanie",         "Éléa Théron",                         "elea.theron@laregion.fr",                              "01/04/2026", "09/04/2026", "—",          "En attente"),
    ("La Région Occitanie",         "Andra Viglietti",                     "andra.viglietti@laregion.fr",                          "09/04/2026", "—",          "—",          "⚠️ BOUNCE — adresse invalide"),
    ("Montpellier Métropole",       "Andra Viglietti",                     "andra.viglietti@montpellier.fr",                       "09/04/2026", "—",          "—",          "En attente"),
    ("Montpellier Métropole",       "Stéphanie Benazet-Iannone",           "stephanie.benazet-iannone@montpellier.fr",             "09/04/2026", "—",          "—",          "En attente"),
    ("Mairie de Montpellier",       "Xavier Deraulin",                     "xavier.deraulin@montpellier.fr",                       "10/04/2026", "—",          "—",          "En attente"),
    ("Journalistes indép.",         "Anna Leduigou / Kenza",               "annaleduigou@live.fr / kenza.gdx@gmail.com",            "09/04/2026", "—",          "—",          "En attente"),
]

for i, row in enumerate(france_data, 3):
    bg = GREY_ROW if i % 2 == 0 else None
    for j, val in enumerate(row, 1):
        statut_bg = None
        if j == 7:
            if "✅" in val:  statut_bg = GREEN_OK
            elif "⚠️" in val: statut_bg = RED_WARN
            elif "📵" in val: statut_bg = ORANGE
        data_cell(ws_fr, i, j, val, statut_bg or bg)
    ws_fr.row_dimensions[i].height = 18

# Largeurs colonnes France
widths_fr = [28, 28, 50, 12, 13, 12, 45]
for i, w in enumerate(widths_fr, 1):
    ws_fr.column_dimensions[get_column_letter(i)].width = w

# ═══════════════════════════════════════════════════════════════════════════════
# FEUILLE ESPAGNE
# ═══════════════════════════════════════════════════════════════════════════════
ws_es = wb.create_sheet("🇪🇸 Espagne")
ws_es.freeze_panes = "A3"

ws_es.merge_cells("A1:G1")
t2 = ws_es["A1"]
t2.value = "SUIVI COMMUNIQUÉ DE PRESSE — ISTÒRIA LAB — ESPAGNE"
t2.font = Font(bold=True, color=HEADER_FG, size=13)
t2.fill = PatternFill("solid", fgColor=RED_ES)
t2.alignment = Alignment(horizontal="center", vertical="center")
ws_es.row_dimensions[1].height = 28

for i, col in enumerate(cols, 1):
    header_cell(ws_es, 2, i, col, YELLOW_ES, fg_color="000000")
ws_es.row_dimensions[2].height = 22

espagne_data = [
    ("La Vanguardia",               "Astrid Meseguer",              "Astrid@lavanguardia.es / cultura@lavanguardia.es",       "01/04/2026", "10/04/2026", "—",          "En attente"),
    ("El Periódico",                "Nando Salva",                  "nsalva@elperiodico.es",                                 "01/04/2026", "—",          "—",          "En attente"),
    ("El Periódico",                "Leticia Blanco",               "LBlanco@elperiodico.com",                               "02/04/2026", "—",          "—",          "En attente"),
    ("RTVE — De Película (R1 & R5)","Yolanda Flores Remón",         "yolanda.flores@rtve.es",                                "02/04/2026", "—",          "—",          "✅ CONFIRMÉ — interview le 23/04"),
    ("RTVE",                        "Rufino Sanchez",               "rufino.sanchez@rtve.es",                                "02/04/2026", "10/04/2026", "—",          "✅ Répondu — transféré à Gerardo Sánchez & Javier Sales"),
    ("Kinotico",                    "Iñaki Mayora / Matias Rebolledo","inaki.mayora@kinotico.es / matias.rebolledo@kinotico.es","02/04/2026","10/04/2026","—",           "En attente"),
    ("El Diario",                   "Elena Cabrera / Javier Zurro", "ecabrera@eldiario.es / jzurro@eldiario.es",             "02/04/2026", "10/04/2026", "—",          "En attente"),
    ("3Cat / TV3",                  "Montse Pujol Cullere",         "Cultura@3cat.cat / mpujol.c@3cat.cat",                  "02/04/2026", "10/04/2026", "—",          "En attente"),
    ("Caimán Cuadernos de Cine",    "Jara",                         "caiman.cdc@caimanediciones.es",                         "02/04/2026", "07/04/2026", "—",          "En attente"),
    ("Audiovisual 451",             "Irene / David",                "irene@audiovisual451.com / david@audiovisual451.com",   "02/04/2026", "—",          "—",          "En attente"),
    ("35 mm",                       "Carlos / Alvaro",              "redaccion@35milimetros.es",                             "02/04/2026", "—",          "—",          "📰 PUBLIÉ — 35milimetros.es"),
    ("Cineconn",                    "Arturo / José",                "gestion@cineconn.es",                                   "02/04/2026", "07/04/2026", "—",          "En attente"),
    ("Contracultural",              "Maria Arnau",                  "info@contracultural.es",                                "02/04/2026", "07/04/2026", "—",          "En attente"),
    ("Cortos de Metraje",           "Alejandro",                    "comunicacion@cortosdemetraje.com",                      "02/04/2026", "—",          "—",          "En attente"),
    ("ACN (Agència Catalana)",      "Nicolas / Javier",             "ntomas-lanchon@acn.cat / jrubio@acn.cat",               "02/04/2026", "—",          "—",          "En attente"),
    ("Europa Press",                "Luis Arriola / Alberto Ortega","luisarriola@europapress.es / albertoortega@europapress.es","02/04/2026","—",         "—",          "En attente"),
    ("Cine Invisible",              "Carlos",                       "cineinvisible2010@gmail.com",                           "02/04/2026", "10/04/2026", "—",          "En attente"),
    ("José Manuel Romero (indép.)", "José Manuel",                  "josemromeroser@gmail.com",                              "02/04/2026", "—",          "—",          "En attente"),
]

for i, row in enumerate(espagne_data, 3):
    bg = GREY_ROW if i % 2 == 0 else None
    for j, val in enumerate(row, 1):
        statut_bg = None
        if j == 7:
            if "✅" in val:   statut_bg = GREEN_OK
            elif "📰" in val: statut_bg = BLUE_LIGHT
            elif "⚠️" in val: statut_bg = RED_WARN
        data_cell(ws_es, i, j, val, statut_bg or bg)
    ws_es.row_dimensions[i].height = 18

widths_es = [28, 28, 50, 12, 13, 12, 50]
for i, w in enumerate(widths_es, 1):
    ws_es.column_dimensions[get_column_letter(i)].width = w

# ═══════════════════════════════════════════════════════════════════════════════
# FEUILLE LÉGENDE
# ═══════════════════════════════════════════════════════════════════════════════
ws_leg = wb.create_sheet("Légende")
legende = [
    ("Symbole", "Signification"),
    ("✅",      "Répondu positivement / Confirmé"),
    ("📰",      "Article publié"),
    ("📵",      "Absent / Réponse automatique"),
    ("⚠️",      "Problème (bounce, adresse invalide…)"),
    ("—",       "Pas encore fait"),
    ("En attente", "Aucune réponse reçue"),
]
for i, (sym, sig) in enumerate(legende, 1):
    bold = (i == 1)
    bg = RED_FR if i == 1 else None
    fg = "FFFFFF" if i == 1 else None
    c1 = ws_leg.cell(row=i, column=1, value=sym)
    c2 = ws_leg.cell(row=i, column=2, value=sig)
    for c in (c1, c2):
        c.font = Font(bold=bold, color=fg or "000000", size=10)
        if bg: c.fill = PatternFill("solid", fgColor=bg)
        c.alignment = Alignment(vertical="center")
        c.border = make_border()
ws_leg.column_dimensions["A"].width = 14
ws_leg.column_dimensions["B"].width = 42

# Sauvegarde
out = "/home/user/claude-workspace/suivi_CP_Istoria_Lab.xlsx"
wb.save(out)
print(f"Fichier créé : {out}")
